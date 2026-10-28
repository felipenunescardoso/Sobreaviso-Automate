from datetime import date, datetime, timedelta
from pathlib import Path


from inicio import conn


MESES_PT = {
    1: "janeiro",
    2: "fevereiro",
    3: "março",
    4: "abril",
    5: "maio",
    6: "junho",
    7: "julho",
    8: "agosto",
    9: "setembro",
    10: "outubro",
    11: "novembro",
    12: "dezembro",
}

DIAS_SEMANA_PT = {
    0: "segunda-feira",
    1: "terça-feira",
    2: "quarta-feira",
    3: "quinta-feira",
    4: "sexta-feira",
    5: "sábado",
    6: "domingo",
}


# ============================================================
# COMPETÊNCIA
# ============================================================

def calcular_competencia(data_referencia=None):
    """
    A competência do Sobreaviso sempre é:
        dia 11 do mês anterior até dia 10 do mês da referência.
    """

    if data_referencia is None:
        data_referencia = date.today()

    if isinstance(data_referencia, datetime):
        data_referencia = data_referencia.date()

    primeiro_dia_mes = data_referencia.replace(day=1)
    ultimo_dia_mes_anterior = primeiro_dia_mes - timedelta(days=1)

    competencia_inicio = date(
        ultimo_dia_mes_anterior.year,
        ultimo_dia_mes_anterior.month,
        11,
    )

    competencia_fim = date(
        data_referencia.year,
        data_referencia.month,
        10,
    )

    return competencia_inicio, competencia_fim


# ============================================================
# CONSULTAS
# ============================================================

def buscar_lancamentos(id_usuario, data_referencia=None):

    competencia_inicio, competencia_fim = calcular_competencia(
        data_referencia
    )

    fim_exclusivo = competencia_fim + timedelta(days=1)

    sql = """
        SELECT
            u.id_usuario,
            u.nome,
            u.chapa,
            d.nome_departamento,
            e.id_escala,
            ec.id_escala_competencia,
            l.id_lancamento,
            l.tipo,
            l.data_hora_inicio,
            l.data_hora_fim,
            l.motivo
        FROM sobreaviso.lancamentos l

        JOIN sobreaviso.escala_competencias ec
            ON ec.id_escala_competencia =
               l.id_escala_competencia

        JOIN sobreaviso.escalas e
            ON e.id_escala = ec.id_escala

        JOIN sobreaviso.usuarios u
            ON u.id_usuario = e.id_usuario

        JOIN sobreaviso.departamentos d
            ON d.id_departamento = u.id_departamento

        WHERE u.id_usuario = %s
          AND l.data_hora_inicio >= %s
          AND l.data_hora_inicio < %s

        ORDER BY
            l.data_hora_inicio,
            l.id_lancamento;
    """

    with conn.cursor() as cur:

        cur.execute(
            sql,
            (
                id_usuario,
                competencia_inicio,
                fim_exclusivo
            )
        )

        colunas = [
            coluna.name
            for coluna in cur.description
        ]

        linhas = cur.fetchall()

        registros = [
            dict(zip(colunas, linha))
            for linha in linhas
        ]

    return competencia_inicio, competencia_fim, registros


def buscar_usuario(id_usuario):

    sql = """
        SELECT
            u.id_usuario,
            u.nome,
            u.chapa,
            u.email,
            d.nome_departamento,
            u.id_empresa,
            e.nome_empresa,
            e.banco_de_horas,
            u.id_cidade,
            u.uf
        FROM sobreaviso.usuarios u

        INNER JOIN sobreaviso.departamentos d
            ON d.id_departamento = u.id_departamento

        INNER JOIN sobreaviso.empresas e
            ON e.id_empresa = u.id_empresa

        WHERE u.id_usuario = %s
          AND u.ativo = TRUE;
    """

    with conn.cursor() as cur:

        cur.execute(
            sql,
            (id_usuario,)
        )

        linha = cur.fetchone()

        if not linha:
            return None

        colunas = [
            coluna.name
            for coluna in cur.description
        ]

        return dict(
            zip(colunas, linha)
        )


# ============================================================
# FORMATAÇÃO / CÁLCULOS
# ============================================================

def formatar_duracao(inicio, fim):

    total_segundos = max(
        0,
        int(
            (fim - inicio).total_seconds()
        )
    )

    horas, resto = divmod(
        total_segundos,
        3600
    )

    minutos = resto // 60

    return f"{horas:02d}:{minutos:02d}"


def formatar_segundos(total_segundos):

    total_segundos = max(
        0,
        int(total_segundos)
    )

    horas = total_segundos // 3600
    minutos = (
        total_segundos % 3600
    ) // 60

    return f"{horas:02d}:{minutos:02d}"


def eh_feriado(
    data_referencia,
    uf=None,
    id_cidade=None
):

    with conn.cursor() as cur:

        cur.execute(
            """
            SELECT EXISTS
            (
                SELECT 1
                FROM sobreaviso.feriados f
                WHERE f.data_feriado = %s
                  AND f.ativo = TRUE
                  AND
                  (
                        (
                            f.uf IS NULL
                            AND f.id_cidade IS NULL
                        )
                        OR (
                            f.uf = %s
                            AND f.id_cidade IS NULL
                        )
                        OR (
                            f.id_cidade = %s
                        )
                  )
            );
            """,
            (
                data_referencia,
                uf.upper() if uf else None,
                id_cidade
            )
        )

        return bool(
            cur.fetchone()[0]
        )


def calcular_banco_horas(
    registros,
    usuario
):

    if not usuario.get(
        "banco_de_horas"
    ):
        return 0, 0

    id_cidade = usuario.get(
        "id_cidade"
    )

    uf = usuario.get("uf")

    limite_banco_dia = (
        2 * 60 * 60
    )

    banco_por_dia = {}

    segundos_banco = 0
    segundos_pagar = 0

    acionamentos = [
        r
        for r in registros
        if r["tipo"] == "ACIONAMENTO"
    ]

    acionamentos.sort(
        key=lambda r: (
            r["data_hora_inicio"],
            r["id_lancamento"]
        )
    )

    for registro in acionamentos:

        inicio = registro[
            "data_hora_inicio"
        ]

        fim = registro[
            "data_hora_fim"
        ]

        if fim <= inicio:
            continue

        dia = inicio.date()

        total_segundos = int(
            (
                fim - inicio
            ).total_seconds()
        )

        if (
            inicio.weekday() >= 5
            or eh_feriado(
                dia,
                uf,
                id_cidade
            )
        ):

            segundos_pagar += (
                total_segundos
            )

            continue

        inicio_banco = inicio.replace(
            hour=18,
            minute=1,
            second=0,
            microsecond=0
        )

        fim_banco = inicio.replace(
            hour=23,
            minute=59,
            second=59,
            microsecond=999999
        )

        faixa_inicio = max(
            inicio,
            inicio_banco
        )

        faixa_fim = min(
            fim,
            fim_banco
        )

        segundos_elegiveis = 0

        if faixa_fim > faixa_inicio:

            segundos_elegiveis = int(
                (
                    faixa_fim
                    - faixa_inicio
                ).total_seconds()
            )

        segundos_fora_banco = (
            total_segundos
            - segundos_elegiveis
        )

        segundos_pagar += max(
            0,
            segundos_fora_banco
        )

        if segundos_elegiveis <= 0:
            continue

        ja_usado = banco_por_dia.get(
            dia,
            0
        )

        restante = max(
            0,
            limite_banco_dia
            - ja_usado
        )

        segundos_para_banco = min(
            segundos_elegiveis,
            restante
        )

        segundos_excedentes = (
            segundos_elegiveis
            - segundos_para_banco
        )

        segundos_banco += (
            segundos_para_banco
        )

        segundos_pagar += (
            segundos_excedentes
        )

        banco_por_dia[dia] = (
            ja_usado
            + segundos_para_banco
        )

    return (
        segundos_banco,
        segundos_pagar
    )


# ============================================================
# NOME DO ARQUIVO
# ============================================================

def nome_arquivo_excel(
    usuario,
    competencia_fim
):

    nome = (
        usuario["nome"]
        .strip()
        .replace(" ", "_")
    )

    nome = "".join(
        caractere
        for caractere in nome
        if (
            caractere.isalnum()
            or caractere in ("_", "-")
        )
    )

    return (
        f"Sobreaviso_{nome}_"
        f"{competencia_fim.year}_"
        f"{competencia_fim.month:02d}.xlsx"
    )


# ============================================================
# GERAÇÃO DO EXCEL
# ============================================================

def gerar_excel(
    id_usuario,
    data_referencia=None,
    pasta_saida="excel"
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    usuario = buscar_usuario(id_usuario)

    if not usuario:
        raise ValueError(
            f"Usuário {id_usuario} não encontrado ou inativo."
        )

    competencia_inicio, competencia_fim, registros = buscar_lancamentos(
        id_usuario=id_usuario,
        data_referencia=data_referencia
    )

    pasta = Path(pasta_saida)
    pasta.mkdir(parents=True, exist_ok=True)

    caminho_excel = pasta / nome_arquivo_excel(
        usuario,
        competencia_fim
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Sobreaviso"

    borda_fina = Side(style="thin", color="000000")
    borda = Border(
        left=borda_fina,
        right=borda_fina,
        top=borda_fina,
        bottom=borda_fina
    )
    fill_rotulo = PatternFill("solid", fgColor="F2F2F2")
    fill_cabecalho = PatternFill("solid", fgColor="D9E1F2")

    # Título
    ws.merge_cells("A1:J1")
    ws["A1"] = "SOBREAVISO MENSAL T.I."
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # Informações do relatório
    ws["A3"] = "EMPRESA:"
    ws["B3"] = usuario["nome_empresa"]
    ws.merge_cells("B3:E3")
    ws["F3"] = "DEPARTAMENTO:"
    ws["G3"] = usuario["nome_departamento"]
    ws.merge_cells("G3:J3")

    ws["A4"] = "DATA INICIAL:"
    ws["B4"] = competencia_inicio.strftime("%d/%m/%Y")
    ws.merge_cells("B4:E4")
    ws["F4"] = "DATA FINAL:"
    ws["G4"] = competencia_fim.strftime("%d/%m/%Y")
    ws.merge_cells("G4:J4")

    ws["A5"] = "FUNCIONÁRIO:"
    ws["B5"] = usuario["nome"]
    ws.merge_cells("B5:E5")
    ws["F5"] = "CHAPA:"
    ws["G5"] = str(usuario["chapa"] or "")
    ws.merge_cells("G5:J5")

    for linha in range(3, 6):
        for coluna in range(1, 11):
            celula = ws.cell(linha, coluna)
            celula.border = borda
            celula.alignment = Alignment(vertical="center")

    for ref in ("A3", "A4", "A5", "F3", "F4", "F5"):
        ws[ref].fill = fill_rotulo
        ws[ref].font = Font(bold=True)

    # Cabeçalho principal
    cabecalhos = [
        "CHAPA",
        "NOME",
        "DT SOBREAVISO",
        "HORA INICIAL",
        "HORA FINAL",
        "MÊS",
        "OBSERVAÇÕES",
        "TIPO",
        "HORAS TOTAIS",
        "DIA DA SEMANA"
    ]

    for coluna, titulo in enumerate(cabecalhos, start=1):
        celula = ws.cell(7, coluna, titulo)
        celula.fill = fill_cabecalho
        celula.font = Font(bold=True)
        celula.border = borda
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    total_segundos = 0
    total_acionamentos = 0
    linha_atual = 8

    if registros:
        for registro in registros:
            inicio = registro["data_hora_inicio"]
            fim = registro["data_hora_fim"]
            tipo = registro["tipo"]

            total_segundos += max(
                0,
                int((fim - inicio).total_seconds())
            )

            if tipo == "ACIONAMENTO":
                total_acionamentos += 1

            valores = [
                str(usuario["chapa"] or ""),
                usuario["nome"],
                inicio.strftime("%d/%m/%Y"),
                inicio.strftime("%H:%M"),
                fim.strftime("%H:%M"),
                MESES_PT[inicio.month],
                registro["motivo"] or "",
                tipo,
                formatar_duracao(inicio, fim),
                DIAS_SEMANA_PT[inicio.weekday()]
            ]

            for coluna, valor in enumerate(valores, start=1):
                celula = ws.cell(linha_atual, coluna, valor)
                celula.border = borda
                celula.alignment = Alignment(
                    horizontal=("left" if coluna in (2, 7) else "center"),
                    vertical="center",
                    wrap_text=(coluna == 7)
                )

            linha_atual += 1
    else:
        ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=10)
        ws["A8"] = "Nenhum lançamento encontrado para esta competência."
        ws["A8"].alignment = Alignment(horizontal="center", vertical="center")
        ws["A8"].border = borda
        linha_atual = 9

    # Resumo
    linha_resumo = linha_atual + 1
    horas_totais = total_segundos // 3600
    minutos_totais = (total_segundos % 3600) // 60

    if usuario["banco_de_horas"]:
        segundos_banco, segundos_pagar = calcular_banco_horas(
            registros,
            usuario
        )
        segundos_sobreaviso = max(
            0,
            total_segundos - segundos_banco - segundos_pagar
        )

        resumo = [
            "Quantidade de acionamentos:", total_acionamentos,
            "Total de horas lançadas:", f"{horas_totais:02d}:{minutos_totais:02d}",
            "Banco de horas:", formatar_segundos(segundos_banco),
            "Horas a pagar:", formatar_segundos(segundos_pagar),
            "Horas sobreaviso:", formatar_segundos(segundos_sobreaviso)
        ]
    else:
        resumo = [
            "Quantidade de acionamentos:", total_acionamentos,
            "Total de horas lançadas:", f"{horas_totais:02d}:{minutos_totais:02d}"
        ]

    for coluna, valor in enumerate(resumo, start=1):
        celula = ws.cell(linha_resumo, coluna, valor)
        celula.border = borda
        celula.alignment = Alignment(vertical="center", wrap_text=True)
        if coluna % 2 == 1:
            celula.fill = fill_rotulo
            celula.font = Font(bold=True)

    # Larguras
    larguras = [14, 28, 16, 14, 14, 14, 38, 18, 16, 20]
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura

    ws.freeze_panes = "A8"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:7"
    ws.print_area = f"A1:J{linha_resumo}"

    wb.save(caminho_excel)

    return caminho_excel


# ============================================================
# FUNÇÕES DE APOIO DO FRONT DE EXTRAÇÃO
# ============================================================

MESES_EXTRACAO = [
    "Janeiro",
    "Fevereiro",
    "Março",
    "Abril",
    "Maio",
    "Junho",
    "Julho",
    "Agosto",
    "Setembro",
    "Outubro",
    "Novembro",
    "Dezembro",
]


def listar_tecnicos_extracao():

    with conn.cursor() as cur:

        cur.execute("""
            SELECT
                id_usuario,
                nome
            FROM sobreaviso.usuarios
            WHERE ativo = TRUE
              AND perfil = 'TECNICO'
            ORDER BY nome;
        """)

        return cur.fetchall()


def listar_competencias_extracao():

    competencias = {}

    for mes_fim in range(
        1,
        13
    ):

        mes_inicio = (
            12
            if mes_fim == 1
            else mes_fim - 1
        )

        nome = (
            f"{MESES_EXTRACAO[mes_inicio - 1]}"
            f" - "
            f"{MESES_EXTRACAO[mes_fim - 1]}"
        )

        competencias[
            nome
        ] = mes_fim

    return competencias


def montar_data_referencia_competencia(
    ano,
    mes_fim
):

    return date(
        int(ano),
        int(mes_fim),
        14
    )


def consultar_extracao(
    id_usuario,
    ano,
    mes_fim
):

    data_referencia = (
        montar_data_referencia_competencia(
            ano=ano,
            mes_fim=mes_fim
        )
    )

    (
        competencia_inicio,
        competencia_fim,
        registros
    ) = buscar_lancamentos(
        id_usuario=id_usuario,
        data_referencia=data_referencia
    )

    return {
        "id_usuario": id_usuario,
        "data_referencia":
            data_referencia,
        "competencia_inicio":
            competencia_inicio,
        "competencia_fim":
            competencia_fim,
        "registros":
            registros
    }


# ============================================================
# PERMISSÕES DA EXTRAÇÃO
# ============================================================

def _buscar_perfil_extracao(
    id_usuario_executor
):

    with conn.cursor() as cur:

        cur.execute("""
            SELECT perfil
            FROM sobreaviso.usuarios
            WHERE id_usuario = %s
              AND ativo = TRUE;
        """, (
            id_usuario_executor,
        ))

        resultado = cur.fetchone()

    if not resultado:

        raise ValueError(
            "Usuário executor "
            "não encontrado ou inativo."
        )

    return str(
        resultado[0]
    ).upper()


def _ids_subordinados_extracao(
    id_usuario_gestor
):

    with conn.cursor() as cur:

        cur.execute("""
            WITH RECURSIVE subordinados AS
            (
                SELECT
                    u.id_usuario,
                    u.id_gestor
                FROM sobreaviso.usuarios u
                WHERE u.id_gestor = %s
                  AND u.ativo = TRUE

                UNION ALL

                SELECT
                    u.id_usuario,
                    u.id_gestor
                FROM sobreaviso.usuarios u

                INNER JOIN subordinados s
                    ON u.id_gestor =
                       s.id_usuario

                WHERE u.ativo = TRUE
            )

            SELECT DISTINCT
                id_usuario
            FROM subordinados;
        """, (
            id_usuario_gestor,
        ))

        return {
            linha[0]
            for linha in cur.fetchall()
        }


def validar_permissao_extracao(
    id_usuario_executor,
    id_usuario_alvo
):

    perfil = (
        _buscar_perfil_extracao(
            id_usuario_executor
        )
    )

    if perfil == "ADMIN":
        return True

    if perfil == "TECNICO":

        if (
            id_usuario_executor
            != id_usuario_alvo
        ):

            raise PermissionError(
                "Técnico sem permissão "
                "para extrair Excel "
                "de outro usuário."
            )

        return True

    if perfil == "GESTOR":

        subordinados = (
            _ids_subordinados_extracao(
                id_usuario_executor
            )
        )

        if (
            id_usuario_alvo
            not in subordinados
        ):

            raise PermissionError(
                "Gestor sem permissão "
                "para extrair Excel "
                "deste usuário."
            )

        return True

    raise PermissionError(
        "Perfil sem permissão "
        f"para extração: {perfil}"
    )


def listar_usuarios_permitidos_extracao(
    id_usuario_executor
):

    perfil = (
        _buscar_perfil_extracao(
            id_usuario_executor
        )
    )

    with conn.cursor() as cur:

        if perfil == "ADMIN":

            cur.execute("""
                SELECT
                    id_usuario,
                    nome
                FROM sobreaviso.usuarios
                WHERE ativo = TRUE
                  AND perfil = 'TECNICO'
                ORDER BY nome;
            """)

            return cur.fetchall()

        if perfil == "GESTOR":

            cur.execute("""
                WITH RECURSIVE subordinados AS
                (
                    SELECT
                        u.id_usuario,
                        u.nome,
                        u.perfil,
                        u.id_gestor
                    FROM sobreaviso.usuarios u
                    WHERE u.id_gestor = %s
                      AND u.ativo = TRUE

                    UNION ALL

                    SELECT
                        u.id_usuario,
                        u.nome,
                        u.perfil,
                        u.id_gestor
                    FROM sobreaviso.usuarios u

                    INNER JOIN subordinados s
                        ON u.id_gestor =
                           s.id_usuario

                    WHERE u.ativo = TRUE
                )

                SELECT DISTINCT
                    id_usuario,
                    nome
                FROM subordinados
                WHERE perfil = 'TECNICO'
                ORDER BY nome;
            """, (
                id_usuario_executor,
            ))

            return cur.fetchall()

        if perfil == "TECNICO":

            cur.execute("""
                SELECT
                    id_usuario,
                    nome
                FROM sobreaviso.usuarios
                WHERE id_usuario = %s
                  AND ativo = TRUE;
            """, (
                id_usuario_executor,
            ))

            return cur.fetchall()

    return []


def consultar_extracao_permitida(
    id_usuario_executor,
    id_usuario_alvo,
    ano,
    mes_fim
):

    validar_permissao_extracao(
        id_usuario_executor,
        id_usuario_alvo
    )

    return consultar_extracao(
        id_usuario=id_usuario_alvo,
        ano=ano,
        mes_fim=mes_fim
    )


def gerar_excel_permitido(
    id_usuario_executor,
    id_usuario_alvo,
    data_referencia=None,
    pasta_saida="excel"
):

    validar_permissao_extracao(
        id_usuario_executor,
        id_usuario_alvo
    )

    return gerar_excel(
        id_usuario=id_usuario_alvo,
        data_referencia=data_referencia,
        pasta_saida=pasta_saida
    )


# ============================================================
# TESTE MANUAL
# ============================================================

def ler_data_referencia(valor):

    valor = valor.strip()

    if not valor:
        return date.today()

    return datetime.strptime(
        valor,
        "%d/%m/%Y"
    ).date()


if __name__ == "__main__":

    print(
        "=== EXTRATOR EXCEL "
        "DE SOBREAVISO ==="
    )

    id_usuario = int(
        input(
            "ID do usuário: "
        ).strip()
    )

    referencia = input(
        "Data de referência "
        "[DD/MM/AAAA] "
        "(Enter = hoje): "
    )

    data_referencia = (
        ler_data_referencia(
            referencia
        )
    )

    (
        competencia_inicio,
        competencia_fim
    ) = calcular_competencia(
        data_referencia
    )

    print(
        "Competência selecionada: "
        f"{competencia_inicio:%d/%m/%Y} "
        f"até "
        f"{competencia_fim:%d/%m/%Y}"
    )

    arquivo = gerar_excel(
        id_usuario=id_usuario,
        data_referencia=
            data_referencia
    )

    print(
        f"Excel gerado: {arquivo}")